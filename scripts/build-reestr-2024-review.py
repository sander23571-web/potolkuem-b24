#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-reestr-2024-review.py — 15.09.2026.

Новый файл владельца: «Выписка_Александр_Жуков_2024_год_АДАТ.xlsx» (весь 2024 год,
789 расходных + 7 доходных строк, 8-колоночный формат — как «...с 21 июля 2026.xlsx»,
см. import-reestr-platezhey-2026-09-03.py). Записей за 2024 год в СП «Реестр платежей»
(1080) сейчас нет вообще — самые ранние текущие записи с мая 2025.

По аналогии с обработкой «Выписка 2025 Александр Жуков.xlsx» (03.09.2026): доходы и
дивиденды отделяются от расходов, расходы классифицируются тем же классификатором.
В отличие от 03.09 — здесь используется САМАЯ СВЕЖАЯ версия классификатора (v2 из
audit-reestr-platezhey-2026-09-05.py: фиксы regex от 06.09 + исключение ЗП/премий
отдельным контуром, решение руководства 06.09.2026), не старая v1/classify2 из 09-03.

15.09.2026, доработка по запросу владельца: вместо голого regex-классификатора сначала
пробуем сопоставить платёж по КОНТРАГЕНТУ с уже вручную выверенными живыми записями
СП «Реестр платежей» (2226 шт. после применения v4-аудита 09.09.2026, см. commit ca8ecce)
— эти коды прошли ручную сверку владельца/руководства, они надёжнее регулярки на
незнакомых 2024-формулировках. Правило: если контрагент встречается в живых записях
CRM и ВСЕ его платежи там имеют ОДИН и тот же код — берём этот код напрямую (флаг
"по контрагенту (2025-2026)"). Если у контрагента коды расходятся (напр. «Технопарк
Элма» — 3 разных кода в зависимости от назначения платежа) — не гадаем, отдаём
regex/classify_v2 (для «Технопарка» там уже есть спец-логика по тексту назначения).
Проверка (58 общих контрагентов, из них только 4 неоднозначных) — покрывает ~70% строк
2024-файла точным историческим кодом контрагента.

НЕ пишет ничего в Б24 — только офлайн-классификация и xlsx для проверки владельцем,
импорт в CRM (crm.item.add) делаем отдельным шагом после подтверждения (см. feedback
"пробная партия перед массовой записью").

Запуск: python3 scripts/build-reestr-2024-review.py
"""
import importlib.util
import json
import re
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

PROXY = "https://b24proxy.bobp.ru/b24/potolkuem"
API_KEY = "b589caa6cef8e95163dc9ded06b0934479023f15ef9ab13a76368a74b9694f1c"

SCRIPT_DIR = Path(__file__).parent
BASE_DIR = SCRIPT_DIR.parent
SRC_FILE = BASE_DIR / "Выписка_Александр_Жуков_2024_год_АДАТ.xlsx"
OUT_FILE = BASE_DIR / "Реестр-платежей-2024-к-разбору-2026-09-15.xlsx"

DIVIDEND_RE = re.compile(r"дивиденд", re.IGNORECASE)


def b24(method, params):
    req = urllib.request.Request(
        f"{PROXY}/{method}",
        data=json.dumps(params).encode(),
        headers={"X-API-Key": API_KEY, "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=30) as r:
        d = json.loads(r.read())
    if "error" in d:
        raise RuntimeError(f"{method} -> {d}")
    return d["result"]


def load_code_map():
    resp = b24("lists.element.get", {"IBLOCK_TYPE_ID": "lists", "IBLOCK_ID": 32})
    out = {}
    id_to_code = {}
    for el in resp:
        code_prop = el.get("PROPERTY_182") or {}
        code_val = list(code_prop.values())[0] if code_prop else None
        cat_prop = el.get("PROPERTY_184") or {}
        cat_val = list(cat_prop.values())[0] if cat_prop else None
        if code_val:
            out[code_val] = cat_val or el.get("NAME") or ""
            id_to_code[int(el["ID"])] = code_val
    return out, id_to_code


def split_title(title):
    if " — " in (title or ""):
        contragent, purpose = title.split(" — ", 1)
        return contragent.strip(), purpose.strip()
    return "", title or ""


def build_contragent_code_map(id_to_code):
    """Контрагент -> код, только для контрагентов с ЕДИНСТВЕННЫМ кодом среди всех живых
    записей 1080 (уже вручную выверенных v4-аудитом 09.09.2026). Неоднозначные контрагенты
    (несколько разных кодов) сюда не попадают — по ним решает classify_v2 по тексту."""
    items, start = [], 0
    while True:
        resp = b24("crm.item.list", {
            "entityTypeId": 1080, "select": ["id", "title", "ufCrm32Code"], "start": start,
        })
        batch = resp["items"]
        items.extend(batch)
        if len(batch) < 50:
            break
        start += 50

    from collections import defaultdict, Counter
    per_contragent = defaultdict(Counter)
    for it in items:
        contragent, _ = split_title(it.get("title"))
        cid = it.get("ufCrm32Code")
        if not contragent or not cid:
            continue
        code = id_to_code.get(int(cid))
        if code:
            per_contragent[contragent.lower().strip()][code] += 1

    unanimous = {c: next(iter(cnt)) for c, cnt in per_contragent.items() if len(cnt) == 1}
    ambiguous = {c: cnt for c, cnt in per_contragent.items() if len(cnt) > 1}
    print(f"Контрагентов в живом CRM (1080): {len(per_contragent)}, "
          f"однозначных: {len(unanimous)}, неоднозначных (решает classify_v2): {len(ambiguous)}")
    return unanimous


# ── Переиспользуем самый свежий классификатор (v2, 06.09.2026 фиксы) из audit-скрипта ──
_spec = importlib.util.spec_from_file_location(
    "audit_mod", SCRIPT_DIR / "audit-reestr-platezhey-2026-09-05.py")
_audit = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_audit)
classify_v2 = _audit.classify_v2


def ru_date_to_iso(s):
    d, m, y = str(s).split(".")
    return f"{y}-{m}-{d}"


def classify_hybrid(purpose, contragent, contragent_map):
    """Приоритет: (1) ЗП/премии — regex по PAYROLL_PATTERNS (жёсткое правило, не переопределяется
    историей контрагента); (2) контрагент однозначно встречается в живых записях 1080 — берём
    его исторический код; (3) иначе classify_v2 (regex по тексту назначения)."""
    combined = f"{contragent or ''} {purpose or ''}".lower()
    for pat in _audit.PAYROLL_PATTERNS:
        if re.search(pat, combined, re.IGNORECASE):
            return "ИСКЛЮЧИТЬ-ЗП", "regex-ЗП"
    key = (contragent or "").lower().strip()
    if key in contragent_map:
        return contragent_map[key], "контрагент (2025-2026)"
    return classify_v2(purpose, contragent), "regex (classify_v2)"


def load_file():
    wb = openpyxl.load_workbook(SRC_FILE, data_only=True)
    ws = wb.active
    expenses, income = [], []
    # тот же формат колонок, что и «...с 21 июля 2026.xlsx»: Дата, №док, Дебет, Кредит, Контрагент, ИНН, Назначение
    for r in ws.iter_rows(min_row=13, values_only=True):
        date = r[0]
        if date is None:
            continue
        debet, kredit = r[2], r[3]
        row = {"date": date, "docnum": r[1], "contragent": r[4], "inn": r[5], "purpose": r[6]}
        if debet is not None:
            expenses.append({**row, "amount": debet})
        elif kredit is not None:
            income.append({**row, "amount": kredit})
    return expenses, income


def main():
    print(f"Читаю {SRC_FILE.name}...")
    expenses, income = load_file()
    print(f"Расходов (Дебет): {len(expenses)}, доходов (Кредит, не расход — отдельный лист): {len(income)}")

    dividends = [r for r in expenses if DIVIDEND_RE.search(r["purpose"] or "")]
    expenses = [r for r in expenses if not DIVIDEND_RE.search(r["purpose"] or "")]
    if dividends:
        print(f"Дивиденды исключены из расходов: {len(dividends)} шт.")

    print("Загружаю справочник кодов (Универсальный список, IBLOCK_ID=32)...")
    code_map, id_to_code = load_code_map()

    print("Строю карту контрагент -> код по живым записям СП «Реестр платежей» (1080)...")
    contragent_map = build_contragent_code_map(id_to_code)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "К разбору"
    headers = ["Дата", "№ документа", "Контрагент", "ИНН", "Назначение платежа",
               "Сумма", "Код", "Категория", "Источник кода", "Флаг"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws_excl = wb.create_sheet("Исключено (ЗП, премии)")
    ws_excl.append(["Дата", "Контрагент", "Назначение платежа", "Сумма"])
    for cell in ws_excl[1]:
        cell.font = Font(bold=True)

    ws_inc = wb.create_sheet("Доходы (не расход)")
    ws_inc.append(["Дата", "Контрагент", "Назначение платежа", "Сумма"])
    for cell in ws_inc[1]:
        cell.font = Font(bold=True)

    ws_div = wb.create_sheet("Дивиденды (не расход)")
    ws_div.append(["Дата", "Контрагент", "Назначение платежа", "Сумма"])
    for cell in ws_div[1]:
        cell.font = Font(bold=True)

    fb_fill = PatternFill("solid", fgColor="FFF3CD")
    stats = {"ok": 0, "payroll": 0, "code138": 0, "by_contragent": 0, "by_regex": 0, "total_amount": 0}

    for r in expenses:
        code, source = classify_hybrid(r["purpose"], r["contragent"], contragent_map)
        if code == "ИСКЛЮЧИТЬ-ЗП":
            stats["payroll"] += 1
            ws_excl.append([r["date"], r["contragent"], r["purpose"], r["amount"]])
            continue
        stats["ok"] += 1
        stats["total_amount"] += r["amount"]
        if source.startswith("контрагент"):
            stats["by_contragent"] += 1
        else:
            stats["by_regex"] += 1
        flag = ""
        if code == "13.8":
            stats["code138"] += 1
            flag = "? Прочее — проверить вручную"
        category = code_map.get(code, "")
        ws.append([r["date"], r["docnum"], r["contragent"], r["inn"], r["purpose"],
                   r["amount"], code, category, source, flag])
        if flag:
            for cell in ws[ws.max_row]:
                cell.fill = fb_fill

    for r in income:
        ws_inc.append([r["date"], r["contragent"], r["purpose"], r["amount"]])
    for r in dividends:
        ws_div.append([r["date"], r["contragent"], r["purpose"], r["amount"]])

    widths = {"К разбору": [12, 14, 32, 16, 55, 12, 8, 28, 24, 28],
              "Исключено (ЗП, премии)": [12, 32, 55, 12],
              "Доходы (не расход)": [12, 32, 55, 12],
              "Дивиденды (не расход)": [12, 32, 55, 12]}
    for sheet_name, ws_widths in widths.items():
        sh = wb[sheet_name]
        for i, w in enumerate(ws_widths, start=1):
            sh.column_dimensions[chr(64 + i)].width = w

    wb.save(OUT_FILE)
    print(f"\nСохранено: {OUT_FILE.relative_to(BASE_DIR)}")
    print("Статистика:")
    print(f"  К разбору (в реестр): {stats['ok']} шт., {stats['total_amount']:,.0f} ₽")
    print(f"    из них по коду контрагента (2025-2026, вручную выверено): {stats['by_contragent']} шт.")
    print(f"    из них по regex/classify_v2: {stats['by_regex']} шт.")
    print(f"  Из них в 13.8 «Прочее»: {stats['code138']} ({stats['code138']/stats['ok']*100:.1f}%)")
    print(f"  Исключено как ЗП/премии/отпускные/больничные: {stats['payroll']} шт.")
    print(f"  Доходы (Кредит, не пишем в 1080): {len(income)} шт.")
    print(f"  Дивиденды (не расход): {len(dividends)} шт.")


if __name__ == "__main__":
    main()
