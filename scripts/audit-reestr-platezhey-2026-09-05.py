#!/usr/bin/env python3
"""
audit-reestr-platezhey-2026-09-05.py — офлайн-аудит СП «Реестр платежей» (1080).

НЕ пишет ничего обратно в Б24. Выгружает все живые записи, прогоняет через
(v1) неизменённый текущий классификатор и (v2) классификатор с точечными
исправлениями найденных багов, сравнивает с кодом, который реально стоит в CRM
сейчас, и размечает статус (оставить / исключить как ЗП / расхождение на проверку).

Результат — xlsx, ничего не применяется автоматически.
"""
import importlib.util
import json
import re
import urllib.request

import openpyxl
from openpyxl.styles import Font, PatternFill

PROXY = "https://b24proxy.bobp.ru/b24/potolkuem"
API_KEY = "b589caa6cef8e95163dc9ded06b0934479023f15ef9ab13a76368a74b9694f1c"


def b24(method, params):
    req = urllib.request.Request(
        f"{PROXY}/{method}",
        data=json.dumps(params).encode(),
        headers={"X-API-Key": API_KEY, "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read())


# ── 1. Загрузить справочник кодов (id элемента списка -> "11.7", имя) ───────
def load_code_map():
    resp = b24("lists.element.get", {"IBLOCK_TYPE_ID": "lists", "IBLOCK_ID": 32})
    out = {}
    for el in resp.get("result", []):
        code_prop = el.get("PROPERTY_182") or {}
        code_val = list(code_prop.values())[0] if code_prop else None
        out[int(el["ID"])] = {"code": code_val, "name": el["NAME"]}
    return out


# ── 2. Подключить v1 — неизменённый текущий классификатор ───────────────────
spec = importlib.util.spec_from_file_location(
    "imp0829", "scripts/import-reestr-platezhey-2026-08-29.py"
)
mod_v1 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod_v1)
classify_v1 = mod_v1.classify

# ── 3. v2 — точечные исправления найденных багов ────────────────────────────
PAYROLL_PATTERNS = [
    r"заработн", r"зарплат", r"\bпреми\w*\s+за\b", r"отпу.?кн", r"больничн",
    r"расчёт с персоналом", r"выходное пособие", r"окончательный расчет",
    r"окончательный расчёт", r"подотчет", r"подотчёт",
]

CONTRAGENT_RULES_V2 = list(mod_v1.CONTRAGENT_RULES)
# долфинс -> не 10.2 и не свой код, а 10.1 (КУБ) — Долфинс делает разработку КУБ
CONTRAGENT_RULES_V2 = [
    (p, "10.1") if p == r"долфинс.{0,5}айти" else (p, c)
    for p, c in CONTRAGENT_RULES_V2
]
# подтверждённые вручную владельцем/руководством 03-04.09.2026 — эталонные примеры
CONTRAGENT_RULES_V2 = [
    (r"красикова елена", "12.2"),
    (r"дмитриев кирилл", "12.2"),
    (r"еоц.?партнёр|еоц.?партнер", "6.1"),
    (r"рамазанова эльвира", "13.3"),
    (r"буравлёва юлия|буравлева юлия", "13.3"),
    (r"элкод", "6.2"),  # финальная сверка 06.09.2026: лицензия справочно-правовой системы, не юрист
] + CONTRAGENT_RULES_V2

# точечные правки по конкретному ID из финальной сверки владельца — вынесены во
# внешний версионируемый файл (scripts/reestr-platezhey-overrides.json), не в код,
# чтобы повторный прогон не расходился в зависимости от того, кто и когда его читает
def load_id_overrides():
    try:
        with open("scripts/reestr-platezhey-overrides.json", encoding="utf-8") as f:
            raw = json.load(f)
        return {int(k): v for k, v in raw.items()}
    except FileNotFoundError:
        return {}


def load_id_excludes():
    try:
        with open("scripts/reestr-platezhey-exclude-ids.json", encoding="utf-8") as f:
            return set(json.load(f))
    except FileNotFoundError:
        return set()


def load_id_questions():
    try:
        with open("scripts/reestr-platezhey-question-ids.json", encoding="utf-8") as f:
            return set(json.load(f))
    except FileNotFoundError:
        return set()


ID_OVERRIDES = load_id_overrides()
ID_EXCLUDES = load_id_excludes()
ID_QUESTIONS = load_id_questions()

RULES_V2 = list(mod_v1.RULES)
FIXED_RULES_PREPEND = [
    # более специфичные правила — ПЕРЕД общими (первое совпадение побеждает):
    # отчуждение исключительного авторского права — это юридическая сделка (4.1),
    # не работа дизайнера/иллюстратора (13.3), даже если платёж тому же человеку
    # и даже если рядом упоминается слово "прототип" как объект передачи прав
    (r"отчуждени.{0,15}(исключительн\w*\s+)?авторск\w*\s+прав", "4.1"),
    # беглая гласная: "образец" (не только "образцы/образцов")
    (r"сигнальн.{0,15}образ[ец]{0,2}ц?", "13.7"),
    # "иллюстрации"/"иллюстрация", не только "иллюстратор"
    (r"иллюстра", "13.3"),
    # прототип без обязательного соседства слова "игр"
    (r"прототип", "13.3"),
    # копирайтинг/статьи для сайта -> новый код 11.7 "Журнал" (подтверждено 06.09.2026)
    (r"копирайтинг|стать[ияей].{0,10}(для\s+)?сайт", "11.7"),
    # отдельный проект "Мировая литература" — весь пул платежей по нему это 13.3
    (r"мировая\s+литератур", "13.3"),
    # отдельный проект "Соль Земли" — тот же случай, весь пул это 13.3
    (r"соль\s+земли", "13.3"),
    # мультфильм/сценарий Козырева — оставить в 13.8, разберутся отдельно (не трогать)
]


def classify_v2(purpose, contragent="", record_id=None):
    if record_id in ID_OVERRIDES:
        return ID_OVERRIDES[record_id]
    combined = f"{contragent or ''} {purpose or ''}".lower()
    for pat in PAYROLL_PATTERNS:
        if re.search(pat, combined, re.IGNORECASE):
            return "ИСКЛЮЧИТЬ-ЗП"
    for pattern, code in FIXED_RULES_PREPEND:
        if re.search(pattern, combined, re.IGNORECASE):
            return code
    # переиспользуем логику v1, но с CONTRAGENT_RULES_V2
    purpose_l = (purpose or "").lower()
    for pattern, code in CONTRAGENT_RULES_V2:
        if re.search(pattern, (contragent or "").lower(), re.IGNORECASE):
            if pattern.startswith(r"технопарк"):
                if re.search(r"уборк|клининг", purpose_l):
                    return "1.3"
                if re.search(r"автостоянк|въезд.{0,15}автотранспорт|парковк", purpose_l):
                    return "1.4"
                if re.search(r"счетчик|электроэнерг|водоснабж|возмещени", purpose_l):
                    return "1.2"
                return "1.1"
            return code
    for pattern, code in RULES_V2:
        if re.search(pattern, combined, re.IGNORECASE):
            return code
    return "13.8"


# ── 4. Выгрузка всех живых записей ──────────────────────────────────────────
def fetch_all():
    items, start = [], 0
    while True:
        resp = b24("crm.item.list", {
            "entityTypeId": 1080,
            "select": ["id", "title", "opportunity", "begindate", "ufCrm32Code"],
            "start": start,
        })
        batch = resp.get("result", {}).get("items", [])
        items.extend(batch)
        if len(batch) < 50:
            break
        start += 50
    return items


def split_title(title):
    if " — " in (title or ""):
        contragent, purpose = title.split(" — ", 1)
        return contragent.strip(), purpose.strip()
    return "", title or ""


def main():
    print("Загружаю справочник кодов...")
    code_map = load_code_map()

    print("Выгружаю все записи СП «Реестр платежей»...")
    items = fetch_all()
    print(f"Всего записей: {len(items)}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Аудит"
    headers = ["ID", "Контрагент", "Назначение платежа", "Сумма", "Дата",
               "Код (сейчас в CRM)", "Код_v1 (чистый прогон)", "Код_v2 (с фиксами)",
               "Статус", "Комментарий", "Флаг для руководителя"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws_excl = wb.create_sheet("Исключено (ЗП)")
    ws_excl.append(["ID", "Контрагент", "Назначение платежа", "Сумма", "Дата"])
    for cell in ws_excl[1]:
        cell.font = Font(bold=True)

    mismatch_v1_fill = PatternFill("solid", fgColor="FFF3CD")
    question_fill = PatternFill("solid", fgColor="D6E4FF")
    stats = {"match": 0, "mismatch_v1_vs_current": 0, "exclude_payroll": 0, "code138_v2": 0, "flagged": 0}

    for it in items:
        contragent, purpose = split_title(it.get("title"))
        cur_code_id = it.get("ufCrm32Code")
        cur_code = code_map.get(int(cur_code_id), {}).get("code") if cur_code_id else None

        v1 = classify_v1(purpose, contragent)
        v2 = classify_v2(purpose, contragent, record_id=it["id"])

        if v2 == "ИСКЛЮЧИТЬ-ЗП" or it["id"] in ID_EXCLUDES:
            stats["exclude_payroll"] += 1
            ws_excl.append([it["id"], contragent, purpose, it.get("opportunity"), it.get("begindate")])
            continue

        status = "Оставить"
        comment = ""
        if cur_code and v1 != cur_code:
            status = "Расхождение v1 vs текущий код — проверить"
            comment = f"текущий={cur_code}, чистый прогон v1={v1}"
            stats["mismatch_v1_vs_current"] += 1
        else:
            stats["match"] += 1

        if v2 == "13.8":
            stats["code138_v2"] += 1

        is_flagged = it["id"] in ID_QUESTIONS
        if is_flagged:
            stats["flagged"] += 1

        row = [
            it["id"], contragent, purpose, it.get("opportunity"), it.get("begindate"),
            cur_code, v1, v2, status, comment,
            "? — требует решения руководителя" if is_flagged else "",
        ]
        ws.append(row)
        if is_flagged:
            for cell in ws[ws.max_row]:
                cell.fill = question_fill
        elif status.startswith("Расхождение"):
            for cell in ws[ws.max_row]:
                cell.fill = mismatch_v1_fill

    for i, w in enumerate([8, 32, 55, 12, 12, 16, 16, 16, 32, 40, 32], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for i, w in enumerate([8, 32, 55, 12, 12], start=1):
        ws_excl.column_dimensions[chr(64 + i)].width = w

    out_path = "Реестр-платежей-аудит-2026-09-05.xlsx"
    wb.save(out_path)
    print(f"\nСохранено: {out_path}")
    print("Статистика:")
    print(f"  Совпадает с текущим (оставить как есть): {stats['match']}")
    print(f"  К исключению (ЗП/премии/отпускные/больничные): {stats['exclude_payroll']}")
    print(f"  Расхождение v1 vs текущий код в CRM: {stats['mismatch_v1_vs_current']}")
    print(f"  Всё ещё в 13.8 после фиксов v2: {stats['code138_v2']}")
    print(f"  Помечено флагом для руководителя (было под '?'): {stats['flagged']}")


if __name__ == "__main__":
    main()
