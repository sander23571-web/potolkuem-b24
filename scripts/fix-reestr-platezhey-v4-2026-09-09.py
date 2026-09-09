#!/usr/bin/env python3
"""
Применение правок кодов из "Реестр-платежей-аудит-2026-09-05 v4.xlsx" (лист "Аудит") к живым
записям СП «Реестр платежей» (CRM_32, entityTypeId=1080).

Правило (подтверждено владельцем 09.09.2026):
  final_code = E, если E заполнено (столбец "новая ручная правка владельца")
             = "без кода" (очистить поле), если E пусто И L ("Флаг для руководителя") активен
             = I ("Код_v2 с фиксами"), если и E, и L пусты

Долфинс АйТи (10.2→10.1) — уже покрывается общим правилом, отдельной логики не требует.

Режимы: --dry-run (по умолчанию), --pilot (первые 10), --apply (все 339).
"""
import json
import sys
import time
import urllib.request

import openpyxl

PROXY = "https://b24proxy.bobp.ru/b24/potolkuem"
API_KEY = "b589caa6cef8e95163dc9ded06b0934479023f15ef9ab13a76368a74b9694f1c"
ENTITY_TYPE_ID = 1080
XLSX_PATH = "/root/projects/talk/Реестр-платежей-аудит-2026-09-05 v4.xlsx"

CODE_TO_ID = {
    "1.1":162,"1.2":164,"1.3":166,"1.4":168,"1.5":170,"1.6":172,"1.7":174,"1.8":176,
    "1.9":178,"1.10":180,"1.11":182,"1.12":184,
    "2.1":186,"2.2":188,"2.3":190,"3.1":192,"4.1":194,"5.1":196,"6.1":198,"6.2":200,
    "6.3":202,"7.1":204,"7.2":206,"8.1":208,"9.1":210,"10.1":212,"10.2":214,
    "11.1":216,"11.2":218,"11.3":220,"11.4":222,"11.5":224,"11.6":226,
    "12.1":228,"12.2":230,"12.3":232,"12.4":234,"12.5":236,
    "13.1":238,"13.2":240,"13.3":242,"13.4":244,"13.5":246,"13.6":248,"13.7":250,"13.8":252,
    "11.7":254,"12.6":256,
}
CODE_CATALOG = {v: k for k, v in CODE_TO_ID.items()}


def norm(code):
    if code is None:
        return None
    s = str(code).strip()
    s = s.replace("..", ".")
    return s or None


def b24(method, params):
    req = urllib.request.Request(
        f"{PROXY}/{method}",
        data=json.dumps(params).encode(),
        headers={"X-API-Key": API_KEY, "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read())


def fetch_all(entity_type_id, filter_=None, select=None, retries=3):
    select = select or ["*"]
    items = []
    start = 0
    while True:
        for attempt in range(retries):
            try:
                res = b24("crm.item.list", {"entityTypeId": entity_type_id, "filter": filter_ or {}, "select": select, "start": start})
                break
            except Exception as e:
                if attempt == retries - 1:
                    raise
                print(f"  (повтор после ошибки: {e})")
                time.sleep(2)
        batch = res.get("result", {}).get("items", [])
        items.extend(batch)
        if len(batch) < 50:
            break
        start += 50
        time.sleep(0.2)
    return items


def build_plan():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Аудит"]

    rows = []
    for r in range(2, ws.max_row + 1):
        aid = ws.cell(r, 1).value
        if aid is None:
            continue
        e = norm(ws.cell(r, 5).value)
        i = norm(ws.cell(r, 9).value)
        flag = ws.cell(r, 12).value
        if e:
            final = e
        elif flag:
            final = None  # "без кода"
        else:
            final = i
        rows.append((int(aid), final))

    live_items = fetch_all(ENTITY_TYPE_ID, {}, ["id", "ufCrm32Code"])
    live = {}
    for it in live_items:
        c = it.get("ufCrm32Code")
        live[it["id"]] = CODE_CATALOG.get(int(c)) if c else None

    plan = []
    for aid, final in rows:
        if aid not in live:
            print(f"  ПРОПУСК id={aid}: не найден в живом 1080")
            continue
        if final is not None and final not in CODE_TO_ID:
            print(f"  ПРОПУСК id={aid}: код '{final}' не найден в каталоге")
            continue
        if final != live[aid]:
            plan.append((aid, final))
    return plan


def main():
    mode = "dry-run"
    if "--pilot" in sys.argv:
        mode = "pilot"
    elif "--apply" in sys.argv:
        mode = "apply"

    plan = build_plan()
    plan.sort()
    print(f"Всего изменений в плане: {len(plan)}")
    print(f"Режим: {mode}")
    print()

    if mode == "dry-run":
        for aid, final in plan[:20]:
            print(f"  id={aid:<6} -> {'без кода (очистить)' if final is None else final}")
        print(f"  ... и ещё {max(0, len(plan)-20)}")
        return

    batch = plan[:10] if mode == "pilot" else plan
    print(f"Применяю {len(batch)} правок...")
    ok, fail = 0, 0
    results = []
    for aid, final in batch:
        elem_id = CODE_TO_ID[final] if final is not None else 0  # 0 = очистить поле
        res = b24("crm.item.update", {
            "entityTypeId": ENTITY_TYPE_ID,
            "id": aid,
            "fields": {"ufCrm32Code": elem_id},
        })
        if res.get("result"):
            ok += 1
            results.append((aid, final, "OK"))
        else:
            fail += 1
            results.append((aid, final, f"ERROR: {res}"))
            print(f"  ОШИБКА id={aid}: {res}")
        time.sleep(0.15)

    print(f"\nГотово: {ok} успешно, {fail} ошибок")
    with open("scripts/.fix-v4-results.json", "w") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
