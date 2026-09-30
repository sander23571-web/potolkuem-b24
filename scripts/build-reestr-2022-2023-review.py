#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-reestr-2022-2023-review.py — 30.09.2026.

Новые файлы владельца: «Выписка_Александр_Жуков_2022_год_АДАТ.xlsx» (290 строк, 7 колонок —
без № документа и без КПП) и «Выписка_Александр_Жуков_2023_год_АДАТ.xlsx» (380 строк, 8 колонок —
без № документа, но с КПП). Оба формата ОТЛИЧАЮТСЯ от формата 2024-файла (7-8 колонок с
№документа) и от формата 2025/2026-файлов — колонки читаются по отдельной карте на каждый год
(см. YEAR_CONFIG). Записей за 2022/2023 год в СП «Реестр платежей» (1080) сейчас нет вообще —
самые ранние текущие записи с мая 2025.

Классификация — тот же гибрид, что и build-reestr-2024-review.py (15.09.2026): контрагент,
однозначно совпадающий по коду со всеми его живыми записями 1080 за 2025-2026, побеждает;
иначе classify_v2 (regex, audit-reestr-platezhey-2026-09-05.py).

⚠️ ВАЖНОЕ ОТЛИЧИЕ от 2024/2025/2026: наивный DIVIDEND_RE (r"дивиденд") ЗДЕСЬ НЕ ПРИМЕНЯЕТСЯ.
В 2023-файле все 11 строк со словом «дивиденды» — это «Единый налоговый платеж (НДФЛ ...
дивиденды)», т.е. уплата НДФЛ/налога с уже выплаченных дивидендов, а не сама выплата
дивидендов (которой в обоих файлах нет вообще — 0 строк с «перечислени... дивиденд»). Это
законный налоговый платёж компании (уже есть код 2.1 «ЕНП»), исключать его как «не расход»
было бы ошибкой. Такие строки классифицируются как обычные расходы (обычно 2.1), но помечаются
отдельным флагом «ЕНП по дивидендам — проверить» для видимости при ручном разборе. Разделены:
доходы (Кредит) — отдельным листом; настоящих строк-выплат дивидендов в этих файлах не найдено.

НЕ пишет ничего в Б24 — только офлайн-классификация и xlsx для проверки владельцем, импорт
в CRM отдельным шагом после подтверждения (см. feedback "пробная партия перед массовой записью").

Запуск: python3 scripts/build-reestr-2022-2023-review.py
"""
import importlib.util
import json
import re
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

PROXY = "https://b24proxy.bobp.ru/b24/potolkuem"
API_KEY = "b589caa6cef8e95163dc9ded06b0934479023f15ef9ab13a76368a74b9694f1c"

SCRIPT_DIR = Path(__file__).parent
BASE_DIR = SCRIPT_DIR.parent

# год -> (файл, индекс Дата, Дебет, Кредит, Контрагент, ИНН, Назначение, первая строка данных)
YEAR_CONFIG = {
    "2022": {
        "src": BASE_DIR / "Выписка_Александр_Жуков_2022_год_АДАТ.xlsx",
        "out": BASE_DIR / "Реестр-платежей-2022-к-разбору-2026-09-30.xlsx",
        "idx": {"date": 0, "debet": 1, "kredit": 2, "contragent": 3, "inn": 4, "purpose": 5},
        "min_row": 13,
    },
    "2023": {
        "src": BASE_DIR / "Выписка_Александр_Жуков_2023_год_АДАТ.xlsx",
        "out": BASE_DIR / "Реестр-платежей-2023-к-разбору-2026-09-30.xlsx",
        "idx": {"date": 0, "debet": 1, "kredit": 2, "contragent": 3, "inn": 4, "purpose": 6},
        "min_row": 13,
    },
}

ENP_DIVIDEND_FLAG_RE = re.compile(r"дивиденд", re.IGNORECASE)
REAL_DIVIDEND_RE = re.compile(r"перечислен\w*\s+дивиденд|выплат\w*\s+дивиденд", re.IGNORECASE)


def b24(method, params):
    req = urllib.request.Request(
        f"{PROXY}/{method}",
        data=json.dumps(params).encode(),
        headers={"X-API-Key": API_KEY, "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=30) as r:
        d = json.loads(r.read())
    if "error" in d:
        raise RuntimeError(f"{method} -> {d}")
    return d["result"]


def load_code_map():
    resp = b24("lists.element.get", {"IBLOCK_TYPE_ID": "lists", "IBLOCK_ID": 32})
    out, id_to_code = {}, {}
    for el in resp:
        code_prop = el.get("PROPERTY_182") or {}
        code_val = list(code_prop.values())[0] if code_prop else None
        cat_prop = el.get("PROPERTY_184") or {}
        cat_val = list(cat_prop.values())[0] if cat_prop else None
        if code_val:
            out[code_val] = cat_val or el.get("NAME") or ""
            id_to_code[int(el["ID"])] = code_val
    return out, id_to_code


def split_title(title):
    if " — " in (title or ""):
        contragent, purpose = title.split(" — ", 1)
        return contragent.strip(), purpose.strip()
    return "", title or ""


def build_contragent_code_map(id_to_code):
    items, start = [], 0
    while True:
        resp = b24("crm.item.list", {
            "entityTypeId": 1080, "select": ["id", "title", "ufCrm32Code"], "start": start,
        })
        batch = resp["items"]
        items.extend(batch)
        if len(batch) < 50:
            break
        start += 50

    per_contragent = defaultdict(Counter)
    for it in items:
        contragent, _ = split_title(it.get("title"))
        cid = it.get("ufCrm32Code")
        if not contragent or not cid:
            continue
        code = id_to_code.get(int(cid))
        if code:
            per_contragent[contragent.lower().strip()][code] += 1

    unanimous = {c: next(iter(cnt)) for c, cnt in per_contragent.items() if len(cnt) == 1}
    ambiguous = {c: cnt for c, cnt in per_contragent.items() if len(cnt) > 1}
    print(f"Контрагентов в живом CRM (1080): {len(per_contragent)}, "
          f"однозначных: {len(unanimous)}, неоднозначных (решает classify_v2): {len(ambiguous)}")
    return unanimous


# ── переиспользуем самый свежий классификатор (v2, 06.09.2026 фиксы) ────────
_spec = importlib.util.spec_from_file_location(
    "audit_mod", SCRIPT_DIR / "audit-reestr-platezhey-2026-09-05.py")
_audit = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_audit)
classify_v2 = _audit.classify_v2


def classify_hybrid(purpose, contragent, contragent_map):
    combined = f"{contragent or ''} {purpose or ''}".lower()
    for pat in _audit.PAYROLL_PATTERNS:
        if re.search(pat, combined, re.IGNORECASE):
            return "ИСКЛЮЧИТЬ-ЗП", "regex-ЗП"
    key = (contragent or "").lower().strip()
    if key in contragent_map:
        return contragent_map[key], "контрагент (2025-2026)"
    return classify_v2(purpose, contragent), "regex (classify_v2)"


def load_file(cfg):
    wb = openpyxl.load_workbook(cfg["src"], data_only=True)
    ws = wb.active
    idx = cfg["idx"]
    expenses, income = [], []
    for r in ws.iter_rows(min_row=cfg["min_row"], values_only=True):
        date = r[idx["date"]]
        if date is None:
            continue
        debet, kredit = r[idx["debet"]], r[idx["kredit"]]
        row = {
            "date": date, "docnum": None,
            "contragent": r[idx["contragent"]], "inn": r[idx["inn"]],
            "purpose": r[idx["purpose"]],
        }
        if debet is not None:
            expenses.append({**row, "amount": debet})
        elif kredit is not None:
            income.append({**row, "amount": kredit})
    return expenses, income


def process_year(year, cfg, code_map, id_to_code, contragent_map):
    print(f"\n{'=' * 60}\nГОД {year}: {cfg['src'].name}\n{'=' * 60}")
    expenses, income = load_file(cfg)
    print(f"Расходов (Дебет): {len(expenses)}, доходов (Кредит, не расход): {len(income)}")

    real_dividends = [r for r in expenses if REAL_DIVIDEND_RE.search(r["purpose"] or "")]
    expenses = [r for r in expenses if r not in real_dividends]
    if real_dividends:
        print(f"Настоящих выплат дивидендов (исключены из расходов): {len(real_dividends)} шт.")
    else:
        print("Настоящих выплат дивидендов не найдено (0 строк).")

    enp_dividend_mentions = sum(
        1 for r in expenses if ENP_DIVIDEND_FLAG_RE.search(r["purpose"] or "")
    )
    if enp_dividend_mentions:
        print(f"⚠️ Строк с упоминанием «дивиденды», но это ЕНП/налог (НЕ выплата, оставлены "
              f"как расход, помечены флагом): {enp_dividend_mentions} шт.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "К разбору"
    headers = ["Дата", "№ документа", "Контрагент", "ИНН", "Назначение платежа",
               "Сумма", "Код", "Категория", "Источник кода", "Флаг"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws_excl = wb.create_sheet("Исключено (ЗП, премии)")
    ws_excl.append(["Дата", "Контрагент", "Назначение платежа", "Сумма"])
    for cell in ws_excl[1]:
        cell.font = Font(bold=True)

    ws_inc = wb.create_sheet("Доходы (не расход)")
    ws_inc.append(["Дата", "Контрагент", "Назначение платежа", "Сумма"])
    for cell in ws_inc[1]:
        cell.font = Font(bold=True)

    ws_div = wb.create_sheet("Дивиденды (не расход)")
    ws_div.append(["Дата", "Контрагент", "Назначение платежа", "Сумма"])
    for cell in ws_div[1]:
        cell.font = Font(bold=True)

    fb_fill = PatternFill("solid", fgColor="FFF3CD")
    enp_div_fill = PatternFill("solid", fgColor="D6E4FF")
    stats = {"ok": 0, "payroll": 0, "code138": 0, "by_contragent": 0, "by_regex": 0,
             "total_amount": 0, "enp_dividend": 0}

    for r in expenses:
        code, source = classify_hybrid(r["purpose"], r["contragent"], contragent_map)
        if code == "ИСКЛЮЧИТЬ-ЗП":
            stats["payroll"] += 1
            ws_excl.append([r["date"], r["contragent"], r["purpose"], r["amount"]])
            continue
        stats["ok"] += 1
        stats["total_amount"] += r["amount"]
        if source.startswith("контрагент"):
            stats["by_contragent"] += 1
        else:
            stats["by_regex"] += 1

        is_enp_dividend = bool(ENP_DIVIDEND_FLAG_RE.search(r["purpose"] or ""))
        flag = ""
        if code == "13.8":
            stats["code138"] += 1
            flag = "? Прочее — проверить вручную"
        if is_enp_dividend:
            stats["enp_dividend"] += 1
            flag = (flag + " | " if flag else "") + "ЕНП по дивидендам — проверить (не выплата)"

        category = code_map.get(code, "")
        ws.append([r["date"], r["docnum"], r["contragent"], r["inn"], r["purpose"],
                   r["amount"], code, category, source, flag])
        if is_enp_dividend:
            for cell in ws[ws.max_row]:
                cell.fill = enp_div_fill
        elif flag:
            for cell in ws[ws.max_row]:
                cell.fill = fb_fill

    for r in income:
        ws_inc.append([r["date"], r["contragent"], r["purpose"], r["amount"]])
    for r in real_dividends:
        ws_div.append([r["date"], r["contragent"], r["purpose"], r["amount"]])

    widths = {"К разбору": [12, 14, 32, 16, 55, 12, 8, 28, 24, 32],
              "Исключено (ЗП, премии)": [12, 32, 55, 12],
              "Доходы (не расход)": [12, 32, 55, 12],
              "Дивиденды (не расход)": [12, 32, 55, 12]}
    for sheet_name, ws_widths in widths.items():
        sh = wb[sheet_name]
        for i, w in enumerate(ws_widths, start=1):
            sh.column_dimensions[chr(64 + i)].width = w

    wb.save(cfg["out"])
    print(f"\nСохранено: {cfg['out'].relative_to(BASE_DIR)}")
    print("Статистика:")
    print(f"  К разбору (в реестр): {stats['ok']} шт., {stats['total_amount']:,.0f} ₽")
    print(f"    из них по коду контрагента (2025-2026, вручную выверено): {stats['by_contragent']} шт.")
    print(f"    из них по regex/classify_v2: {stats['by_regex']} шт.")
    print(f"  Из них в 13.8 «Прочее»: {stats['code138']} ({stats['code138'] / stats['ok'] * 100:.1f}%)")
    print(f"  Из них ЕНП по дивидендам (флаг, но включены как расход): {stats['enp_dividend']}")
    print(f"  Исключено как ЗП/премии/отпускные/больничные: {stats['payroll']} шт.")
    print(f"  Доходы (Кредит, не пишем в 1080): {len(income)} шт.")
    print(f"  Настоящих выплат дивидендов (не расход): {len(real_dividends)} шт.")
    return stats


def main():
    print("Загружаю справочник кодов (Универсальный список, IBLOCK_ID=32)...")
    code_map, id_to_code = load_code_map()
    print("Строю карту контрагент -> код по живым записям СП «Реестр платежей» (1080)...")
    contragent_map = build_contragent_code_map(id_to_code)

    for year, cfg in YEAR_CONFIG.items():
        process_year(year, cfg, code_map, id_to_code, contragent_map)


if __name__ == "__main__":
    main()
